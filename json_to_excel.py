#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Convertitore JSON -> Excel per nuovi corsi Syllabus.
Legge i file JSON contenenti i dati dei corsi e genera file Excel (.xlsx) formattati professionalmente.
"""

import os
import sys
import json
import argparse
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Forza la codifica dello standard output su UTF-8 in ambienti Windows
if sys.stdout.encoding != 'utf-8':
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')


def convert_json_to_excel(json_path, excel_path, sheet_name='Nuovi Corsi'):
    """
    Converte un file JSON dei corsi in un file Excel con formattazione premium.
    
    :param json_path: Percorso del file JSON di input.
    :param excel_path: Percorso del file Excel di output (.xlsx).
    :param sheet_name: Nome del foglio di lavoro Excel (default: 'Nuovi Corsi').
    :return: True se la conversione ha successo, False altrimenti.
    """
    try:
        # Verifica esistenza file JSON
        if not os.path.exists(json_path):
            print(f"[ERRORE] Il file JSON '{json_path}' non esiste.")
            return False

        # Carica dati JSON
        with open(json_path, "r", encoding="utf-8") as f:
            data = json.load(f)

        if not data:
            print(f"[AVVISO] Il file JSON '{json_path}' è vuoto o non contiene dati validi.")
            return False

        # Converti in DataFrame
        df = pd.DataFrame(data)

        # Evita errori di formula in Excel sostituendo '===' con '◆◆◆' ed eliminando '=' iniziali dai testi
        for col in df.columns:
            if df[col].dtype == object:
                # Sostituisce i divisori '===' con '◆◆◆'
                df[col] = df[col].apply(lambda x: x.replace("===", "◆◆◆") if isinstance(x, str) else x)
                # Rimuove eventuali '=' all'inizio del testo per evitare l'interpretazione come formula
                df[col] = df[col].apply(lambda x: x.lstrip("=") if isinstance(x, str) and x.startswith("=") else x)

        # Assicurati che la directory di destinazione esista
        os.makedirs(os.path.dirname(excel_path), exist_ok=True)

        # Scrittura del file Excel con pandas e openpyxl
        with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name=sheet_name)
            workbook = writer.book
            worksheet = writer.sheets[sheet_name]

            # Stili
            font_family = 'Segoe UI'
            header_font = Font(name=font_family, size=11, bold=True, color='FFFFFF')
            header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')  # Blu Navy
            body_font = Font(name=font_family, size=10)
            
            thin_border = Border(
                left=Side(style='thin', color='D9D9D9'),
                right=Side(style='thin', color='D9D9D9'),
                top=Side(style='thin', color='D9D9D9'),
                bottom=Side(style='thin', color='D9D9D9')
            )

            # Formattazione dell'intestazione
            for col_idx, col_name in enumerate(df.columns, 1):
                cell = worksheet.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.border = thin_border

            # Imposta altezza riga intestazione
            worksheet.row_dimensions[1].height = 28

            # Formattazione del corpo della tabella
            for row_idx in range(2, len(df) + 2):
                worksheet.row_dimensions[row_idx].height = 40  # Altezza ideale per testi con a capo
                for col_idx in range(1, len(df.columns) + 1):
                    cell = worksheet.cell(row=row_idx, column=col_idx)
                    cell.font = body_font
                    cell.border = thin_border

                    col_name = df.columns[col_idx - 1]
                    
                    # Allineamento e formattazioni specifiche in base al tipo di colonna
                    if col_name in ["Durata", "Livelli di Padronanza", "Ambito"]:
                        cell.alignment = Alignment(horizontal='center', vertical='top', wrap_text=True)
                    elif col_name == "Link Dettagli":
                        url = cell.value
                        if url and str(url).startswith("http"):
                            cell.hyperlink = url
                            cell.font = Font(name=font_family, size=10, color='0563C1', underline='single')  # Colore link classico
                        cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
                    else:
                        cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)

            # Regolazione dinamica delle larghezze delle colonne
            for col in worksheet.columns:
                col_letter = get_column_letter(col[0].column)
                col_name = col[0].value
                
                # Calcola la lunghezza massima stimata delle righe nella colonna
                max_len = len(str(col_name or ''))
                for cell in col[1:]:
                    val_str = str(cell.value or '')
                    # Se ci sono ritorni a capo, considera la riga più lunga all'interno della cella
                    lines = val_str.split('\n')
                    line_len = max(len(l) for l in lines) if lines else 0
                    if line_len > max_len:
                        max_len = line_len

                # Applica limiti alle larghezze per rendere il layout compatto e leggibile
                width = max(max_len + 3, 12)
                if col_name in ["Descrizione Breve", "Descrizione Estesa", "Dettaglio Percorso (Syllabus)"]:
                    width = min(width, 45)  # Evita colonne eccessivamente larghe per testi lunghi
                else:
                    width = min(width, 30)

                worksheet.column_dimensions[col_letter].width = width

            # Abilita le griglie del foglio
            worksheet.views.sheetView[0].showGridLines = True

        print(f"  [OK] Conversione completata con successo: '{excel_path}'")
        return True

    except Exception as e:
        print(f"  [ERRORE] Si è verificato un errore durante la conversione in Excel: {e}")
        return False

def generate_full_catalog_excel(history_path=None, output_excel_path=None):
    """
    Legge lo storico dei corsi (corsi_storico.json) e genera un file Excel con l'intero catalogo.
    
    :param history_path: Percorso del file JSON con lo storico (default: data/corsi_storico.json)
    :param output_excel_path: Percorso del file Excel di output (default: data/excel/catalogo_completo_syllabus.xlsx)
    :return: Percorso del file Excel generato se riuscito, None altrimenti.
    """
    if history_path is None:
        history_path = os.path.join("data", "corsi_storico.json")
    if output_excel_path is None:
        output_excel_path = os.path.join("data", "excel", "catalogo_completo_syllabus.xlsx")

    if not os.path.exists(history_path):
        print(f"[ERRORE] Il file storico '{history_path}' non esiste.")
        return None

    try:
        with open(history_path, "r", encoding="utf-8") as f:
            history_data = json.load(f)

        if not history_data:
            print(f"[AVVISO] Lo storico '{history_path}' è vuoto.")
            return None

        formatted_courses = []
        for url, course in history_data.items():
            formatted_courses.append({
                "Titolo": course.get('title', 'N/A'),
                "Ambito": course.get('ambito', 'N/A'),
                "Descrizione Breve": course.get('desc_short', 'N/A'),
                "Descrizione Estesa": course.get('desc_extended', 'N/A'),
                "Link Dettagli": course.get('detail_url', url),
                "Durata": course.get('duration', 'N/A'),
                "Livelli di Padronanza": course.get('levels', 'N/A'),
                "Dettaglio Percorso (Syllabus)": course.get('syllabus_text', 'N/A')
            })

        os.makedirs(os.path.dirname(output_excel_path), exist_ok=True)
        temp_json = output_excel_path + ".temp.json"
        with open(temp_json, "w", encoding="utf-8") as f:
            json.dump(formatted_courses, f, ensure_ascii=False, indent=2)

        success = convert_json_to_excel(temp_json, output_excel_path, sheet_name="Catalogo Completo")
        if os.path.exists(temp_json):
            os.remove(temp_json)

        if success:
            return output_excel_path
        return None
    except Exception as e:
        print(f"[ERRORE] Errore durante la generazione del catalogo completo in Excel: {e}")
        return None


def main():
    parser = argparse.ArgumentParser(description="Convertitore JSON -> Excel per nuovi corsi Syllabus.")
    parser.add_argument("json_file", nargs="?", help="Percorso del file JSON da convertire. Se omesso, esegue la conversione di tutti i file JSON mancanti nella cartella data/json/.")
    parser.add_argument("--full-catalog", action="store_true", help="Genera l'Excel del catalogo completo da data/corsi_storico.json")
    args = parser.parse_args()

    if args.full_catalog:
        print("[CONVERSIONE] Generazione del catalogo completo in corso...")
        out = generate_full_catalog_excel()
        if out:
            print(f"[OK] Catalogo completo generato con successo: {out}")
        return

    if args.json_file:
        # Conversione di un singolo file specifico
        json_path = args.json_file
        base_name = os.path.basename(json_path)
        excel_name = base_name.replace(".json", ".xlsx")

        # Mantieni la stessa logica di cartelle
        json_dir_normalized = os.path.normpath(os.path.dirname(json_path)).replace("\\", "/")
        if "data/json" in json_dir_normalized or json_dir_normalized.endswith("data/json"):
            excel_path = os.path.join("data", "excel", excel_name)
        else:
            # Se è altrove, salvalo nella stessa cartella ma con estensione modificata
            excel_path = json_path.replace(".json", ".xlsx")

        print(f"[CONVERSIONE] Avvio conversione di '{json_path}' in '{excel_path}'...")
        convert_json_to_excel(json_path, excel_path)
    else:
        # Conversione automatica di tutti i file mancanti in data/json/
        json_dir = os.path.join("data", "json")
        excel_dir = os.path.join("data", "excel")

        if not os.path.exists(json_dir):
            print(f"[ERRORE] La cartella '{json_dir}' non esiste. Esegui prima lo scraper.")
            return

        files = [f for f in os.listdir(json_dir) if f.endswith(".json")]
        if not files:
            print(f"[INFO] Nessun file JSON trovato in '{json_dir}'.")
            return

        print(f"[SCANSIONE] Ricerca di file JSON in '{json_dir}' non ancora convertiti in Excel...")
        converted_count = 0
        
        for f in files:
            json_path = os.path.join(json_dir, f)
            excel_name = f.replace(".json", ".xlsx")
            excel_path = os.path.join(excel_dir, excel_name)

            if not os.path.exists(excel_path):
                print(f"  -> Rilevato '{f}': conversione in corso...")
                if convert_json_to_excel(json_path, excel_path):
                    converted_count += 1
            else:
                # File già convertito precedentemente
                pass

        print(f"[FINE] Processo completato. File convertiti in questa esecuzione: {converted_count}")


if __name__ == "__main__":
    main()
